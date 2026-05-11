"""Generate email/RATED_CONTACTS_AUDIT.xlsx — sortable/filterable rated-contact audit.

One sheet, one row per rated contact, sorted by rating desc then score desc.
Header row frozen, autofilter on, column widths sized for content. The
"Rating" column is fill-colored using the same palette as the Mailspring
plugin's PersonBand component so a glance shows the visual ramp.

The "Why" column surfaces the CSV's `notes` field — for the rating-1 bucket
this typically reads "auto-rated 1 on 2026-05-07: cluster #29 Newsletters …"
which is exactly what you want to scan.

Usage:
    python email/_tools/dump_rated_contacts_xlsx.py
"""

from __future__ import annotations

import csv
import sqlite3
from collections import defaultdict
from pathlib import Path
from urllib.parse import quote

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
CSV_PATH = ROOT / "contacts_to_rate.csv"
OUT_PATH = ROOT / "RATED_CONTACTS_AUDIT.xlsx"
WAREHOUSE_DB = ROOT / "warehouse.sqlite"

# Same palette as the Mailspring PersonBand plugin component.
RATING_FILL = {
    "9": "C83838",
    "8": "D96D54",
    "7": "D99454",
    "6": "C8B66C",
    "5": "79B08C",
    "4": "8EB6DF",
    "3": "A4C8E8",
    "2": "B8D4E8",
    "1": "D8D8D8",
    "0": "EFEFEF",
}
# Light-rating swatches need dark text; saturated ones get white.
LIGHT_SWATCHES = {"1", "2", "3", "4"}


def _build_sender_to_me_map() -> dict[str, list[tuple[str, int]]]:
    uri = f"file:{quote(str(WAREHOUSE_DB))}?mode=ro"
    con = sqlite3.connect(uri, uri=True)
    con.row_factory = sqlite3.Row
    try:
        rows = con.execute("""
            SELECT LOWER(m.sender_addr) AS sender,
                   LOWER(r.addr)        AS me_addr,
                   COUNT(*)             AS n
            FROM messages m
            JOIN recipients r ON r.message_id = m.id
            WHERE LOWER(r.addr) IN (SELECT email FROM me_addresses)
              AND m.sender_addr IS NOT NULL
              AND m.sender_addr != ''
            GROUP BY LOWER(m.sender_addr), LOWER(r.addr)
        """).fetchall()
    finally:
        con.close()
    out: dict[str, list[tuple[str, int]]] = defaultdict(list)
    for row in rows:
        out[row["sender"]].append((row["me_addr"], int(row["n"])))
    for s in out:
        out[s].sort(key=lambda x: x[1], reverse=True)
    return dict(out)


def _row_emails(row: dict) -> list[str]:
    out: list[str] = []
    primary = (row.get("email") or "").strip().lower()
    if primary:
        out.append(primary)
    other = (row.get("other_emails") or "").strip()
    if other:
        for piece in other.split("|"):
            p = piece.strip().lower()
            if p:
                out.append(p)
    return out


def _aggregate_me_addrs(
    row: dict, sender_to_me: dict[str, list[tuple[str, int]]]
) -> dict[str, int]:
    """Per-row dict: me_addr → count, summed across the row's emails."""
    agg: dict[str, int] = defaultdict(int)
    for em in _row_emails(row):
        for me_addr, n in sender_to_me.get(em, []):
            agg[me_addr] += n
    return dict(agg)


def main() -> None:
    sender_to_me = _build_sender_to_me_map()

    rated_rows: list[dict] = []
    total = 0
    with CSV_PATH.open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            total += 1
            r = (row.get("rating") or "").strip()
            if r == "":
                continue
            rated_rows.append(row)

    # Sort: rating desc, score desc.
    def _key(row):
        try:
            rating = int(row.get("rating") or "0")
        except ValueError:
            rating = 0
        try:
            score = int(row.get("score") or "0")
        except ValueError:
            score = 0
        return (-rating, -score)

    rated_rows.sort(key=_key)

    # Pre-compute per-row me-addr counts and find the column set:
    # only me-addresses that appear at least once across rated contacts,
    # ordered by total cross-row volume desc (busiest inbox leftmost).
    per_row_me: list[dict[str, int]] = []
    addr_total: dict[str, int] = defaultdict(int)
    for row in rated_rows:
        agg = _aggregate_me_addrs(row, sender_to_me)
        per_row_me.append(agg)
        for k, v in agg.items():
            addr_total[k] += v
    me_addr_cols: list[str] = [
        a for a, _ in sorted(addr_total.items(), key=lambda kv: kv[1], reverse=True)
    ]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rated Contacts"

    base_headers = [
        "Rating",
        "Name",
        "Email",
        "Sent (you→them)",
        "Recv (them→you)",
        "First yr",
        "Last yr",
        "Years",
        "Score",
        "Why (notes)",
    ]
    headers = base_headers + me_addr_cols + ["Other emails"]
    ws.append(headers)
    bold = Font(bold=True)
    for c in range(1, len(headers) + 1):
        ws.cell(row=1, column=c).font = bold
        ws.cell(row=1, column=c).alignment = Alignment(
            vertical="center", wrap_text=True, horizontal="left",
        )

    for idx, row in enumerate(rated_rows):
        rating = (row.get("rating") or "").strip()
        try:
            rating_int = int(rating)
        except ValueError:
            rating_int = None
        record = [
            rating_int if rating_int is not None else rating,
            (row.get("name") or "").strip(),
            (row.get("email") or "").strip(),
            int(row.get("sent_to_them") or 0) if (row.get("sent_to_them") or "").strip().lstrip("-").isdigit() else (row.get("sent_to_them") or ""),
            int(row.get("recv_from_them") or 0) if (row.get("recv_from_them") or "").strip().lstrip("-").isdigit() else (row.get("recv_from_them") or ""),
            (row.get("first_year") or "").strip(),
            (row.get("last_year") or "").strip(),
            (row.get("years_active") or "").strip(),
            int(row.get("score") or 0) if (row.get("score") or "").strip().lstrip("-").isdigit() else (row.get("score") or ""),
            (row.get("notes") or "").strip(),
        ]
        # One cell per me-address column. Empty if zero (cleaner than 0).
        agg = per_row_me[idx]
        for ma in me_addr_cols:
            v = agg.get(ma, 0)
            record.append(v if v else None)
        record.append((row.get("other_emails") or "").strip().replace(" | ", ", "))
        ws.append(record)

    # Color the Rating column per the plugin palette.
    for r_idx in range(2, ws.max_row + 1):
        cell = ws.cell(row=r_idx, column=1)
        rating_str = str(cell.value).strip() if cell.value is not None else ""
        fill_hex = RATING_FILL.get(rating_str)
        if fill_hex:
            cell.fill = PatternFill("solid", fgColor=fill_hex)
            cell.font = Font(
                bold=True,
                color="333333" if rating_str in LIGHT_SWATCHES else "FFFFFF",
            )
            cell.alignment = Alignment(horizontal="center", vertical="center")

    # Column widths (chars), tuned for content.
    base_widths = [7, 28, 38, 9, 9, 9, 9, 7, 9, 60]
    me_widths = [9] * len(me_addr_cols)  # narrow numeric columns
    other_widths = [50]
    widths = base_widths + me_widths + other_widths
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Wrap the "Why" + Other-emails columns
    why_col = base_headers.index("Why (notes)") + 1
    other_col = len(headers)
    for r_idx in range(2, ws.max_row + 1):
        for col in (why_col, other_col):
            cell = ws.cell(row=r_idx, column=col)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        # Center-align the per-me-addr count cells
        for col in range(len(base_headers) + 1, len(base_headers) + 1 + len(me_addr_cols)):
            cell = ws.cell(row=r_idx, column=col)
            cell.alignment = Alignment(horizontal="center", vertical="center")

    # Tilt the per-me-address header text 45° so 30+ narrow columns stay readable.
    for col in range(len(base_headers) + 1, len(base_headers) + 1 + len(me_addr_cols)):
        cell = ws.cell(row=1, column=col)
        cell.alignment = Alignment(textRotation=45, vertical="bottom", horizontal="center", wrap_text=False)
    ws.row_dimensions[1].height = 110

    # Freeze header AND first 3 cols (rating, name, email) so they stay visible
    # while scrolling across the wide pivot.
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(OUT_PATH)
    print(f"wrote {OUT_PATH}  ({len(rated_rows):,} rated rows of {total:,} total)")


if __name__ == "__main__":
    main()
